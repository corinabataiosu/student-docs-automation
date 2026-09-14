from pathlib import Path

from docx import Document

from src.document_generator import generate_document

def test_generate_document(tmp_path: Path):
    template_path = tmp_path / "template.docx"
    output_path = tmp_path / "output.docx"

    document = Document()
    document.add_paragraph("Student: {{ first_name }} {{ last_name }}")
    document.add_paragraph("ID: {{ student_id }}")
    document.add_paragraph("Request: {{ request_id }}")
    document.save(template_path)

    student = {
        "first_name": "Maria",
        "last_name": "Popescu",
        "student_id": "ST1001",
        "faculty": "Computer Science",
        "specialization": "Computer Science",
        "year": "3",
        "enrollment_status": "active",
    }

    request = {
        "request_id": "REQ-TEST-0001",
        "request_date": "2024-10-01",
    }

    generate_document(
        template_path=template_path,
        output_path=output_path,
        student=student,
        request=request,
    )

    assert output_path.exists()

    generated_document = Document(output_path)

    text = "\n".join([paragraph.text for paragraph in generated_document.paragraphs])

    assert "Maria Popescu" in text
    assert "ST1001" in text
    assert "REQ-TEST-0001" in text
    assert "{{ first_name }}" not in text