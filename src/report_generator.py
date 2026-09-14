from pathlib import Path
from wsgiref import headers

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

HEADERS = ["Request ID", "Student ID", "Document Type", "Request Date", "Status", "Details"]

def generate_report(report_path: Path, results: list[dict]):

    if report_path.exists():
        workbook = load_workbook(report_path)
        sheet = workbook["Request Report"]
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Request Report"

        sheet.append(HEADERS)

        for cell in sheet[1]:
            cell.font = Font(bold=True)

    for result in results:
        sheet.append([
            result["request_id"],
            result["student_id"],
            result["document_type"],
            result["request_date"],
            result["status"],
            result.get("details", ""),
        ])

    sheet.column_dimensions["A"].width = 15
    sheet.column_dimensions["B"].width = 15
    sheet.column_dimensions["C"].width = 20
    sheet.column_dimensions["D"].width = 15
    sheet.column_dimensions["E"].width = 15
    sheet.column_dimensions["F"].width = 25

    report_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(report_path)